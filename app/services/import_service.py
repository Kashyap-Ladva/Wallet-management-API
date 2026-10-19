from datetime import date, datetime, time, timezone
from decimal import Decimal, InvalidOperation
from hashlib import sha256
from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy.orm import Session

from app.models import MWallet, SWallet, Transaction
from app.services import wallet_service
from app.services.wallet_service import normalize_payment_method

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
REQUIRED_COLUMNS = {
    "amount",
    "type",
    "category",
    "date_created",
    "modeofpayment",
    "description",
}
ALLOWED_CONTENT_TYPES = {
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/octet-stream",
}


def validate_upload(filename: str | None, content_type: str | None, content: bytes) -> str:
    if not filename or not filename.lower().endswith(".xlsx"):
        raise ValueError("Only .xlsx Excel files are supported")
    if content_type and content_type not in ALLOWED_CONTENT_TYPES:
        raise ValueError("Unsupported Excel file content type")
    if not content:
        raise ValueError("The uploaded file is empty")
    if len(content) > MAX_UPLOAD_BYTES:
        raise ValueError("Excel file exceeds the 10 MB upload limit")
    return sha256(content).hexdigest()


def _parse_date(value: object) -> datetime:
    if isinstance(value, datetime):
        result = value
    elif isinstance(value, date):
        result = datetime.combine(value, time.min)
    elif isinstance(value, str):
        try:
            result = datetime.fromisoformat(value.strip())
        except ValueError as exc:
            raise ValueError("date_created must be an Excel date or ISO datetime") from exc
    else:
        raise ValueError("date_created must be an Excel date or ISO datetime")

    if result.tzinfo is not None:
        result = result.astimezone(timezone.utc).replace(tzinfo=None)
    return result


def _read_rows(content: bytes) -> tuple[list[dict], list[dict]]:
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except (BadZipFile, InvalidFileException, OSError) as exc:
        raise ValueError("The uploaded file is not a readable .xlsx workbook") from exc

    try:
        worksheet = workbook.active
        header_cells = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = [
            str(value).strip().lower() if value is not None else ""
            for value in header_cells
        ]
        if not headers or not REQUIRED_COLUMNS.issubset(set(headers)):
            missing = sorted(REQUIRED_COLUMNS.difference(headers))
            raise ValueError(f"Missing required columns: {', '.join(missing)}")
        if len(headers) != len(set(headers)):
            raise ValueError("Excel column names must be unique")

        column_indexes = {name: index for index, name in enumerate(headers) if name}
        rows: list[dict] = []
        errors: list[dict] = []

        for row_number, cells in enumerate(
            worksheet.iter_rows(min_row=2, values_only=True),
            start=2,
        ):
            if not any(value is not None for value in cells):
                continue

            raw = {
                name: cells[index] if index < len(cells) else None
                for name, index in column_indexes.items()
            }
            row_errors: list[str] = []
            try:
                amount = Decimal(str(raw.get("amount", "")).strip())
                if not amount.is_finite() or amount <= 0:
                    raise ValueError("amount must be a finite number greater than zero")
            except (InvalidOperation, ValueError):
                row_errors.append("amount must be a finite number greater than zero")
                amount = Decimal("0")

            transaction_type = str(raw.get("type") or "").strip().lower()
            if transaction_type not in {"income", "expense"}:
                row_errors.append("type must be 'income' or 'expense'")

            category = str(raw.get("category") or "").strip()
            if not category:
                row_errors.append("category is required")

            raw_mode = str(raw.get("modeofpayment") or "").strip()
            if not raw_mode:
                row_errors.append("modeofpayment is required")
                modeofpayment = ""
            else:
                modeofpayment = " ".join(raw_mode.split()).upper()

            try:
                date_created = _parse_date(raw.get("date_created"))
            except ValueError as exc:
                row_errors.append(str(exc))
                date_created = datetime.min

            description = raw.get("description")
            proof = raw.get("proof")
            if description is not None:
                description = str(description).strip() or None
            if proof is not None:
                proof = str(proof).strip() or None

            if row_errors:
                errors.append({"row": row_number, "errors": row_errors})
                continue

            rows.append(
                {
                    "row_number": row_number,
                    "amount": amount,
                    "type": transaction_type,
                    "category": category,
                    "date_created": date_created,
                    "modeofpayment": modeofpayment,
                    "description": description,
                    "proof": proof,
                }
            )
        return rows, errors
    finally:
        workbook.close()


def preview_excel(content: bytes) -> dict:
    rows, errors = _read_rows(content)
    return {
        "valid": not errors and bool(rows),
        "rows": rows,
        "errors": errors,
        "duplicate_policy": "allow",
    }


def preview_excel_for_super_wallet(
    db: Session,
    super_wallet: SWallet,
    content: bytes,
) -> dict:
    rows, errors = _read_rows(content)
    available: dict[str, Decimal] = {}

    for row in rows:
        norm = row["modeofpayment"]
        existing = (
            db.query(MWallet)
            .filter(
                MWallet.super_wallet_id == super_wallet.id,
                MWallet.modeofpayment == norm,
            )
            .first()
        )
        if existing is not None:
            row["mini_wallet_id"] = existing.id
            row["mini_wallet_name"] = existing.name
            if norm not in available:
                available[norm] = existing.balance
        else:
            row["mini_wallet_id"] = None
            row["mini_wallet_name"] = norm
            if norm not in available:
                available[norm] = Decimal("0.00")

        if row["type"] == "expense":
            if row["amount"] > available[norm]:
                errors.append(
                    {
                        "row": row["row_number"],
                        "errors": [
                            "expense exceeds the available mini-wallet balance"
                        ],
                    }
                )
            else:
                available[norm] -= row["amount"]
        else:
            available[norm] += row["amount"]

    return {
        "valid": not errors and bool(rows),
        "rows": rows,
        "errors": errors,
        "duplicate_policy": (
            "Duplicate rows are imported as separate transactions; no automatic "
            "deduplication is performed."
        ),
    }


def import_excel_rows_for_super_wallet(
    db: Session,
    super_wallet: SWallet,
    rows: list[dict],
) -> int:
    resolved_minis: dict[str, MWallet] = {}
    available_balances: dict[str, Decimal] = {}

    for row in rows:
        norm = normalize_payment_method(row["modeofpayment"])
        if norm not in resolved_minis:
            mini = wallet_service.resolve_or_create_mini_wallet(
                db,
                user_id=super_wallet.user_id,
                super_wallet_id=super_wallet.id,
                modeofpayment=norm,
            )
            resolved_minis[norm] = mini
            available_balances[norm] = mini.balance

    transactions: list[Transaction] = []
    for row in rows:
        norm = normalize_payment_method(row["modeofpayment"])
        mini = resolved_minis[norm]
        amount = Decimal(str(row["amount"]))
        if row["type"] == "expense":
            if amount > available_balances[norm]:
                raise ValueError(
                    f"Row {row['row_number']}: expense exceeds the available mini-wallet balance"
                )
            available_balances[norm] -= amount
        else:
            available_balances[norm] += amount

        transaction_data = {
            key: value
            for key, value in row.items()
            if key not in {"row_number", "mini_wallet_id", "mini_wallet_name"}
        }
        transaction_data["modeofpayment"] = norm
        transactions.append(
            Transaction(
                mini_wallet_id=mini.id,
                super_wallet_id=super_wallet.id,
                **transaction_data,
            )
        )

    try:
        db.add_all(transactions)
        db.commit()
    except Exception:
        db.rollback()
        raise
    return len(transactions)


def import_excel_rows(
    db: Session,
    mini_wallet: MWallet,
    rows: list[dict],
) -> int:
    return import_excel_rows_for_super_wallet(db, mini_wallet.super_wallet, rows)


def import_preview_rows_for_super_wallet(
    db: Session,
    super_wallet: SWallet,
    content: bytes,
) -> tuple[int, list[dict]]:
    preview = preview_excel_for_super_wallet(db, super_wallet, content)
    if preview["errors"]:
        return 0, preview["errors"]
    if not preview["rows"]:
        raise ValueError("The workbook contains no transaction rows")
    return import_excel_rows_for_super_wallet(db, super_wallet, preview["rows"]), []


def import_preview_rows(
    db: Session,
    mini_wallet: MWallet,
    content: bytes,
) -> tuple[int, list[dict]]:
    return import_preview_rows_for_super_wallet(db, mini_wallet.super_wallet, content)
